"""Local, deterministic two-file reconciliation. Python 3.10+."""
from __future__ import annotations

import argparse
import csv
import html
import json
import re
import sys
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import BadZipFile


def load(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            records = list(csv.reader(handle, strict=True))
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise ValueError('Multiple worksheets: select a sheet explicitly from ' + ', '.join(book.sheetnames))
            if sheet is not None and sheet not in book.sheetnames:
                raise ValueError('Selected worksheet does not exist.')
            records = []
            for row in book[sheet or book.sheetnames[0]].iter_rows():
                if any(cell.data_type == 'f' for cell in row):
                    raise ValueError('Formula cells are unsupported; export calculated values first.')
                records.append([cell.value for cell in row])
        finally:
            book.close()
    else:
        raise ValueError('Inputs must be UTF-8 CSV or XLSX files.')
    if not records:
        raise ValueError('Input has no header.')
    headers = [str(v) if v is not None else '' for v in records[0]]
    if any(not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('Headers must be nonblank and unique.')
    rows = []
    for record in records[1:]:
        if not record or all(v is None or v == '' for v in record):
            continue
        if len(record) != len(headers):
            raise ValueError('A row has a different number of cells than the header.')
        rows.append(dict(zip(headers, [text(v) for v in record])))
    return headers, rows


def text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def number(value):
    try:
        result = Decimal(value)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def timestamp(value):
    if not re.match(r'^\d{4}-\d{2}-\d{2}(?:$|[T ])', value):
        return None
    try:
        return datetime.fromisoformat(value.replace('Z', '+00:00'))
    except ValueError:
        return None


def infer(values):
    values = [v for v in values if v != '']
    if values and all(number(v) is not None for v in values):
        return 'numeric'
    if values and all(timestamp(v) is not None for v in values):
        return 'date'
    return 'string'


def profile(headers, rows):
    columns = {}
    for field in headers:
        values = [r[field] for r in rows]
        populated = [v for v in values if v != '']
        counts = Counter(populated)
        columns[field] = {'type': infer(values), 'nulls': len(values) - len(populated),
                          'distinct': len(counts),
                          'duplicate_rows': sum(n - 1 for n in counts.values()),
                          'candidate_key': bool(rows) and all(v.strip() for v in values) and len(counts) == len(rows)}
    return {'rows': len(rows), 'columns': columns}


def compare_cell(left, right, kind, numeric_tolerance, date_days):
    if kind == 'numeric' and any(v != '' and number(v) is None for v in (left, right)):
        return 'type_error', None
    if kind == 'date' and any(v != '' and timestamp(v) is None for v in (left, right)):
        return 'type_error', None
    if left == right:
        return 'exact', None
    if left == '' or right == '':
        return 'null_mismatch', None
    if kind == 'numeric':
        a, b = number(left), number(right)
        if a is None or b is None:
            return 'type_error', None
        delta = b - a
        return ('within_tolerance' if abs(delta) <= numeric_tolerance else 'numeric_difference'), str(delta)
    if kind == 'date':
        a, b = timestamp(left), timestamp(right)
        if a is None or b is None or (a.tzinfo is None) != (b.tzinfo is None):
            return 'type_error', None
        delta = Decimal(str((b - a).total_seconds())) / Decimal(86400)
        return ('within_tolerance' if abs(delta) <= date_days else 'date_shift'), str(delta)
    return 'string_difference', None


def reconcile(left, right, key, numeric_tolerance='0.01', date_days='0', types=None):
    lh, lr = left
    rh, rr = right
    tol, days = Decimal(numeric_tolerance), Decimal(date_days)
    if not tol.is_finite() or not days.is_finite() or tol < 0 or days < 0:
        raise ValueError('Tolerances must be finite and nonnegative.')
    for headers, rows in (left, right):
        if key not in headers:
            raise ValueError('Join key must exist on both sides.')
        keys = [r[key] for r in rows]
        if any(not k.strip() for k in keys) or len(set(keys)) != len(keys):
            raise ValueError('Join keys must be nonblank and unique; resolve duplicates upstream.')
    fields = [f for f in lh if f in rh and f != key]
    if not fields:
        raise ValueError('No shared non-key fields to compare.')
    types = types or {}
    if any(f not in fields or t not in ('numeric', 'date', 'string') for f, t in types.items()):
        raise ValueError('Type overrides must name shared fields and numeric, date, or string.')
    kinds = {f: types.get(f, infer([r[f] for r in lr + rr])) for f in fields}
    li, ri = {r[key]: r for r in lr}, {r[key]: r for r in rr}
    shared = li.keys() & ri.keys()
    only_left, only_right = li.keys() - ri.keys(), ri.keys() - li.keys()
    counts = Counter(exact_rows=0, tolerance_rows=0, mismatched_rows=0)
    field_stats = {f: Counter() for f in fields}
    buckets = Counter()
    samples = []
    offsets = {f: Counter() for f in fields}
    for k in sorted(shared):
        statuses = []
        for f in fields:
            status, delta = compare_cell(li[k][f], ri[k][f], kinds[f], tol, days)
            statuses.append(status)
            field_stats[f][status] += 1
            if status != 'exact':
                buckets[status] += 1
            if delta is not None and status in ('numeric_difference', 'date_shift'):
                offsets[f][delta] += 1
            if status not in ('exact', 'within_tolerance') and len(samples) < 50:
                samples.append({'key': k, 'field': f, 'left': li[k][f], 'right': ri[k][f],
                                'status': status, 'delta': delta})
        outcome = ('mismatched_rows' if any(s not in ('exact', 'within_tolerance') for s in statuses)
                   else 'tolerance_rows' if 'within_tolerance' in statuses else 'exact_rows')
        counts[outcome] += 1
    findings = []
    for f, counter in offsets.items():
        if counter:
            delta, n = counter.most_common(1)[0]
            if n >= 2:
                findings.append({'field': f, 'evidence': f'{n} differing cells share target-minus-source delta {delta}',
                                 'hypothesis': 'Possible effective-date shift' if kinds[f] == 'date' else 'Possible systematic offset; rounding or transformation needs verification',
                                 'next_step': 'Check source transformation and effective-date rules before changing tolerances.'})
    if only_left or only_right:
        findings.append({'evidence': f'{len(only_left)} source-only and {len(only_right)} target-only keys',
                         'hypothesis': 'Coverage gap or changed identifiers; cause is unconfirmed',
                         'next_step': 'Compare extraction filters, feed coverage, and key construction.'})
    schema = {'source_only': [f for f in lh if f not in rh], 'target_only': [f for f in rh if f not in lh]}
    total = len(shared) + len(only_left) + len(only_right)
    summary = {'version': 1, 'created_utc': datetime.now(timezone.utc).isoformat(), 'key': key,
               'source_rows': len(lr), 'target_rows': len(rr), 'shared_keys': len(shared),
               **dict(counts), 'source_only_rows': len(only_left), 'target_only_rows': len(only_right),
               'match_rate': (counts['exact_rows'] + counts['tolerance_rows']) / total if total else None,
               'schema_difference': schema, 'types': kinds,
               'rules': {'numeric_absolute_tolerance': str(tol), 'date_tolerance_days': str(days),
                         'nulls': 'empty cells only; two empty cells match', 'strings': 'exact, case and whitespace sensitive'},
               'field_stats': {f: dict(c) for f, c in field_stats.items()},
               'cell_buckets': dict(buckets), 'findings': findings,
               'status': 'differences' if counts['mismatched_rows'] or only_left or only_right or any(schema.values()) else 'tied'}
    detail = {'samples': samples, 'source_only_keys': sorted(only_left)[:50],
              'target_only_keys': sorted(only_right)[:50]}
    return summary, detail


def render(summary, detail):
    esc = lambda v: html.escape(str(v))
    cards = ''.join(f'<article><strong>{summary[k]}</strong><span>{label}</span></article>' for k, label in
                    [('exact_rows', 'Exact rows'), ('tolerance_rows', 'Within tolerance'),
                     ('mismatched_rows', 'Mismatched rows'), ('source_only_rows', 'Source only'), ('target_only_rows', 'Target only')])
    fields = ''.join(f'<tr><td>{esc(f)}</td><td>{esc(summary["types"][f])}</td><td>{esc(json.dumps(s))}</td></tr>' for f, s in summary['field_stats'].items())
    samples = ''.join('<tr>' + ''.join(f'<td>{esc(row[k])}</td>' for k in ('key', 'field', 'left', 'right', 'status', 'delta')) + '</tr>' for row in detail['samples'])
    findings = ''.join(f'<li><b>{esc(f["evidence"])}</b><p>{esc(f["hypothesis"])}</p>{esc(f["next_step"])}</li>' for f in summary['findings']) or '<li>No repeated offset detected. Inspect field counts and extraction rules for remaining differences.</li>'
    bars = ''.join(f'<div class="bar"><span>{esc(k)}: {v}</span><meter min="0" max="{max(summary["cell_buckets"].values(), default=1)}" value="{v}"></meter></div>' for k, v in summary['cell_buckets'].items())
    rate = 'N/A' if summary['match_rate'] is None else f'{summary["match_rate"]:.1%}'
    return f'''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Tie-Out | Reconciliation report</title><style>
body{{font:16px system-ui,sans-serif;background:#f4f6f9;color:#18283b;margin:0}}main{{max-width:1150px;margin:auto;padding:40px 24px}}header{{border-top:5px solid #16817a;padding:24px 0}}h1{{font-size:46px;margin:8px 0}}h2{{margin-top:36px}}.eyebrow{{letter-spacing:3px;color:#16817a;font-size:12px;font-weight:700}}.cards{{display:flex;gap:12px;flex-wrap:wrap}}article{{background:white;border:1px solid #dce3eb;border-radius:12px;padding:20px;flex:1;min-width:120px}}strong{{display:block;font-size:32px}}article span{{color:#536379}}table{{border-collapse:collapse;width:100%;background:white}}td,th{{padding:12px;border-bottom:1px solid #dce3eb;text-align:left;overflow-wrap:anywhere}}.scroll{{overflow:auto}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;background:white;padding:18px;border-radius:8px}}li{{margin:20px 0}}.bar{{display:flex;justify-content:space-between;gap:20px;margin:12px 0}}meter{{width:45%}}footer{{margin-top:40px;color:#536379;font-size:13px}}</style>
<main><header><div class="eyebrow">TIE-OUT / LOCAL DATA RECONCILIATION</div><h1>Do the numbers agree?</h1><p>Status: <b>{esc(summary['status'])}</b> · {rate} matching keys across the union · {summary['shared_keys']} shared keys</p></header>
<section class="cards">{cards}</section><h2>Discrepancy patterns</h2><p>Cell counts, not row counts. Tolerance matches are accepted differences, not errors.</p>{bars or '<p>No differing cells.</p>'}
<h2>Diagnosis & next steps</h2><p>Hypotheses describe possible causes, not proven root causes.</p><ul>{findings}</ul>
<h2>Field comparison</h2><div class="scroll"><table><tr><th>Field</th><th>Type</th><th>Cell counts</th></tr>{fields}</table></div>
<h2>Mismatch examples</h2><p>First 50 differing cells in sorted key order; values are local report contents.</p><div class="scroll"><table><tr><th>Key</th><th>Field</th><th>Source</th><th>Target</th><th>Category</th><th>Delta</th></tr>{samples}</table></div>
<h2>Missing keys</h2><pre>{esc(json.dumps({k:v for k,v in detail.items() if k != 'samples'}, indent=2))}</pre>
<h2>Comparison contract</h2><pre>{esc(json.dumps({'key':summary['key'], **summary['rules'], 'schema_difference':summary['schema_difference']}, indent=2))}</pre>
<footer>Generated {esc(summary['created_utc'])}. Match rate = (exact + within-tolerance rows) / union of keys. Empty inputs have no match rate. Schema differences prevent a tied status. No data was uploaded by this script.</footer></main></html>'''


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode', choices=['profile', 'compare'])
    parser.add_argument('source')
    parser.add_argument('target')
    parser.add_argument('--key')
    parser.add_argument('--source-sheet')
    parser.add_argument('--target-sheet')
    parser.add_argument('--numeric-tolerance', default='0.01')
    parser.add_argument('--date-days', default='0')
    parser.add_argument('--type', action='append', default=[], metavar='COLUMN=TYPE')
    parser.add_argument('--out', default='reports/tie-out')
    parser.add_argument('--fail-on-difference', action='store_true')
    args = parser.parse_args(argv)
    try:
        left, right = load(args.source, args.source_sheet), load(args.target, args.target_sheet)
        if args.mode == 'profile':
            lp, rp = profile(*left), profile(*right)
            candidates = [f for f in left[0] if f in right[0] and lp['columns'][f]['candidate_key'] and rp['columns'][f]['candidate_key']]
            print(json.dumps({'source': lp, 'target': rp, 'candidate_keys': candidates}, indent=2))
            return 0
        if not args.key:
            raise ValueError('Compare requires --key; run profile to inspect candidate keys.')
        overrides = dict(item.split('=', 1) for item in args.type)
        summary, detail = reconcile(left, right, args.key, args.numeric_tolerance, args.date_days, overrides)
        out = Path(args.out)
        targets = [out / 'summary.json', out / 'report.html']
        if any(p.resolve() in (Path(args.source).resolve(), Path(args.target).resolve()) for p in targets):
            raise ValueError('Output must not overwrite an input.')
        if any(p.exists() for p in targets):
            raise ValueError('Output already exists; choose a new --out directory.')
        out.mkdir(parents=True, exist_ok=True)
        targets[0].write_text(json.dumps(summary, indent=2), encoding='utf-8')
        targets[1].write_text(render(summary, detail), encoding='utf-8')
        print(json.dumps({'summary': summary, 'report': str(targets[1].resolve())}, indent=2))
        return 1 if args.fail_on_difference and summary['status'] == 'differences' else 0
    except (ValueError, OSError, csv.Error, ImportError, InvalidOperation, BadZipFile) as error:
        print(f'tie-out: {error}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
