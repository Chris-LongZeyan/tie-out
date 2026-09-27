"""Local, deterministic reconciliation of files and SQL Server queries. Python 3.10+."""
from __future__ import annotations

import argparse
import csv
import html
import heapq
import json
import os
import re
import sys
from collections import Counter
from datetime import date, datetime, timezone
from decimal import Decimal, InvalidOperation, getcontext, localcontext
from difflib import SequenceMatcher
from itertools import chain
from operator import itemgetter
from pathlib import Path
from zipfile import BadZipFile


def load(path, sheet=None):
    path = Path(path)
    if path.suffix.lower() == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as handle:
            return read_records(csv.reader(handle, strict=True), strings=True)
    elif path.suffix.lower() == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            if sheet is None and len(book.sheetnames) != 1:
                raise ValueError('Multiple worksheets: select a sheet explicitly from ' + ', '.join(book.sheetnames))
            if sheet is not None and sheet not in book.sheetnames:
                raise ValueError('Selected worksheet does not exist.')
            def records():
                for row in book[sheet or book.sheetnames[0]].iter_rows():
                    if any(cell.data_type == 'f' for cell in row):
                        raise ValueError('Formula cells are unsupported; export calculated values first.')
                    yield [cell.value for cell in row]
            return read_records(records())
        finally:
            book.close()
    else:
        raise ValueError('Inputs must be UTF-8 CSV or XLSX files.')


def read_records(records, strings=False):
    records = iter(records)
    header = next(records, None)
    if header is None:
        raise ValueError('Input has no header.')
    headers = [str(v) if v is not None else '' for v in header]
    if any(not h.strip() for h in headers) or len(set(headers)) != len(headers):
        raise ValueError('Headers must be nonblank and unique.')
    rows = []
    for record in records:
        blank = not any(record) if strings else all(v is None or v == '' for v in record)
        if blank:
            continue
        if len(record) != len(headers):
            raise ValueError('A row has a different number of cells than the header.')
        rows.append(dict(zip(headers, record if strings else map(text, record))))
    return headers, rows


def text(value):
    if value is None:
        return ''
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return str(value)


def sql_text(value):
    return text(value)


def readonly_query(query):
    tokens = re.findall(r"'(?:''|[^'])*'|\[(?:\]\]|[^\]])*\]|\"(?:\"\"|[^\"])*\"|--[^\r\n]*|/\*.*?\*/|[A-Za-z_]+|[^\s]", query, re.S)
    tokens = [token.upper() for token in tokens if not token.startswith(('--', '/*'))]
    if tokens and tokens[-1] == ';':
        tokens.pop()
    forbidden = {'INSERT', 'UPDATE', 'DELETE', 'MERGE', 'INTO', 'EXEC', 'EXECUTE',
                 'DROP', 'ALTER', 'CREATE', 'TRUNCATE', 'GRANT', 'REVOKE', 'DENY'}
    if not tokens or tokens[0] not in ('SELECT', 'WITH') or ';' in tokens or forbidden.intersection(tokens):
        raise ValueError('Query must be one read-only SELECT or WITH statement; this is a guardrail, not a security boundary.')
    return query.strip()


def load_sql(query, dsn, max_rows, timeout):
    if max_rows <= 0 or timeout <= 0:
        raise ValueError('--max-rows and --timeout must be positive integers.')
    if not dsn:
        raise ValueError('SQL Server input needs --mssql-dsn or the TIEOUT_MSSQL_DSN environment variable.')
    query = readonly_query(query)
    try:
        import pyodbc
    except ImportError:
        raise ValueError('SQL Server input needs pyodbc; install requirements-mssql.txt.')
    connection = cursor = None
    try:
        connection = pyodbc.connect(dsn, timeout=timeout)
        connection.timeout = timeout
        cursor = connection.cursor()
        cursor.execute(query)
        header = [str(d[0]) for d in cursor.description or []]
        if not header or any(not h.strip() for h in header) or len(set(header)) != len(header):
            raise ValueError('Query columns must be named, nonblank, and unique.')
        rows = []
        while True:
            batch = cursor.fetchmany(10000)
            if not batch:
                break
            rows.extend(dict(zip(header, (sql_text(value) for value in row))) for row in batch)
            if len(rows) > max_rows:
                raise ValueError(f'Query returned more rows than the --max-rows limit of {max_rows}.')
        return header, rows
    except pyodbc.Error:
        raise ValueError('SQL Server operation failed; check connection settings, read permissions, query, and timeout locally.') from None
    finally:
        for resource in (cursor, connection):
            if resource is not None:
                try:
                    resource.close()
                except pyodbc.Error:
                    pass


def number(value):
    try:
        result = Decimal(value)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


ISO_PREFIX = re.compile(r'^\d{4}-\d{2}-\d{2}(?:$|[T ])')


def timestamp(value):
    if len(value) < 10 or not ISO_PREFIX.match(value):
        return None
    try:
        return datetime.fromisoformat(value.replace('Z', '+00:00'))
    except ValueError:
        return None


def infer(values):
    populated = False
    numeric, dates = True, True
    seen = set()
    for value in values:
        if value == '' or value in seen:
            continue
        if len(seen) < 4096:
            seen.add(value)
        populated = True
        if numeric and number(value) is not None:
            dates = False
            continue
        numeric = False
        if dates and timestamp(value) is None:
            dates = False
        if not numeric and not dates:
            return 'string'
    return 'numeric' if populated and numeric else 'date' if populated and dates else 'string'


def profile(headers, rows):
    columns = {}
    for field in headers:
        values = [r[field] for r in rows]
        populated = [v for v in values if v != '']
        counts = Counter(populated)
        columns[field] = {'type': infer(values), 'nulls': len(values) - len(populated),
                          'distinct': len(counts),
                          'duplicate_rows': sum(n - 1 for n in counts.values()),
                          'candidate_key': bool(rows) and all(v.strip() for v in values) and len(counts) == len(rows)}
    return {'rows': len(rows), 'columns': columns}


def compare_cell(left, right, kind, numeric_tolerance, date_days):
    if kind in ('numeric', 'date'):
        parse = number if kind == 'numeric' else timestamp
        a = parse(left) if left != '' else None
        b = a if left == right else parse(right) if right != '' else None
        if (left != '' and a is None) or (right != '' and b is None):
            return 'type_error', None
    if left == right:
        return 'exact', None
    if left == '' or right == '':
        return 'null_mismatch', None
    if kind == 'numeric':
        needed = max(a.adjusted(), b.adjusted()) - min(a.as_tuple().exponent, b.as_tuple().exponent) + 2
        if needed > getcontext().prec:
            with localcontext() as context:
                context.prec = max(context.prec, needed)
                delta = b - a
        else:
            delta = b - a
        return ('within_tolerance' if delta.copy_abs() <= numeric_tolerance else 'numeric_difference'), str(delta)
    if kind == 'date':
        if (a.tzinfo is None) != (b.tzinfo is None):
            return 'type_error', None
        elapsed = b - a
        microseconds = (elapsed.days * 86400 + elapsed.seconds) * 1000000 + elapsed.microseconds
        tolerance = Decimal(date_days)
        with localcontext() as context:
            context.prec = max(context.prec, len(tolerance.as_tuple().digits) + 11)
            accepted = abs(microseconds) <= tolerance * 86400000000
            delta = Decimal(microseconds) / Decimal(86400000000)
        return ('within_tolerance' if accepted else 'date_shift'), str(delta)
    return 'string_difference', None


def reconcile(left, right, key, numeric_tolerance='0.01', date_days='0', types=None,
              column_tolerances=None, column_days=None, fuzzy=None):
    lh, lr = left
    rh, rr = right
    keys = [k.strip() for k in key.split(',')] if isinstance(key, str) else [k.strip() for k in key]
    if not keys or any(not k for k in keys) or len(set(keys)) != len(keys):
        raise ValueError('Join key must name one or more distinct columns.')
    tol, days = Decimal(numeric_tolerance), Decimal(date_days)
    if not tol.is_finite() or not days.is_finite() or tol < 0 or days < 0:
        raise ValueError('Tolerances must be finite and nonnegative.')
    column_tolerances, column_days = column_tolerances or {}, column_days or {}
    for mapping, label in ((column_tolerances, 'numeric tolerance'), (column_days, 'date tolerance')):
        for value in mapping.values():
            parsed = Decimal(str(value))
            if not parsed.is_finite() or parsed < 0:
                raise ValueError(f'Per-column {label} values must be finite and nonnegative.')
    if len(keys) == 1:
        k0 = keys[0]
        row_key = itemgetter(k0)
        joined = lambda k: k
    else:
        row_key = lambda r: tuple(r[k] for k in keys)
        joined = lambda k: json.dumps(k, ensure_ascii=False)
    indexes = []
    for headers, rows in (left, right):
        if any(k not in headers for k in keys):
            raise ValueError('Join key must exist on both sides.')
        index = {row_key(r): r for r in rows}
        if len(keys) == 1:
            blank = any(not k.strip() for k in index)
        else:
            blank = any(any(not c.strip() for c in pair) for pair in index)
        if blank or len(index) != len(rows):
            raise ValueError('Join keys must be nonblank and unique; resolve duplicates upstream.')
        indexes.append(index)
    fields = [f for f in lh if f in rh and f not in keys]
    if not fields:
        raise ValueError('No shared non-key fields to compare.')
    types = types or {}
    if any(f not in fields or t not in ('numeric', 'date', 'string') for f, t in types.items()):
        raise ValueError('Type overrides must name shared fields and numeric, date, or string.')
    if any(f not in fields for f in list(column_tolerances) + list(column_days)):
        raise ValueError('Per-column tolerances must name shared non-key fields.')
    kinds = {f: types[f] if f in types else infer(map(itemgetter(f), chain(lr, rr))) for f in fields}
    for mapping, kind in ((column_tolerances, 'numeric'), (column_days, 'date')):
        if any(kinds[f] != kind for f in mapping):
            raise ValueError(f'Per-column {kind} tolerances require {kind} fields; check --type overrides.')
    tols = {f: Decimal(str(column_tolerances[f])) if f in column_tolerances else tol for f in fields}
    dayss = {f: Decimal(str(column_days[f])) if f in column_days else days for f in fields}
    li, ri = indexes
    only_left, only_right = li.keys() - ri.keys(), ri.keys() - li.keys()
    shared_count = len(li) - len(only_left)
    counts = Counter(exact_rows=0, tolerance_rows=0, mismatched_rows=0)
    field_stats = {f: Counter() for f in fields}
    buckets = Counter()
    samples = []
    offsets = {f: Counter() for f in fields}
    total_mismatches = 0
    fast_equal = not any(kind != 'string' for kind in types.values())
    pending = [k for k, row in li.items() if k in ri and row != ri[k]] if fast_equal else li.keys() & ri.keys()
    equal_rows = shared_count - len(pending)
    for k in sorted(pending):
        source_row, target_row = li[k], ri[k]
        statuses = []
        key_text = joined(k)
        for f in fields:
            status, delta = compare_cell(source_row[f], target_row[f], kinds[f], tols[f], dayss[f])
            statuses.append(status)
            field_stats[f][status] += 1
            if status != 'exact':
                buckets[status] += 1
            if delta is not None and status in ('numeric_difference', 'date_shift'):
                offsets[f][Decimal(delta)] += 1
            if status not in ('exact', 'within_tolerance'):
                sample = {'key': key_text, 'field': f, 'left': source_row[f], 'right': target_row[f],
                          'status': status, 'delta': delta}
                numeric = delta is not None and status in ('numeric_difference', 'date_shift')
                rank = (int(numeric), Decimal(delta).copy_abs() if numeric else Decimal(0),
                        -total_mismatches)
                total_mismatches += 1
                heapq.heappush(samples, (rank, sample))
                if len(samples) > 50:
                    heapq.heappop(samples)
        outcome = ('mismatched_rows' if any(s not in ('exact', 'within_tolerance') for s in statuses)
                   else 'tolerance_rows' if 'within_tolerance' in statuses else 'exact_rows')
        counts[outcome] += 1

    counts['exact_rows'] += equal_rows
    if equal_rows:
        for stats in field_stats.values():
            stats['exact'] += equal_rows
    samples = [sample for _, sample in sorted(samples, key=lambda item: item[0], reverse=True)]
    findings = []
    for f, counter in offsets.items():
        if counter:
            delta, n = counter.most_common(1)[0]
            if n >= 2:
                findings.append({'field': f, 'evidence': f'{n} differing cells share target-minus-source delta {delta}',
                                 'hypothesis': 'Possible effective-date shift' if kinds[f] == 'date' else 'Possible systematic offset; rounding or transformation needs verification',
                                 'next_step': 'Check source transformation and effective-date rules before changing tolerances.'})
    if only_left or only_right:
        findings.append({'evidence': f'{len(only_left)} source-only and {len(only_right)} target-only keys',
                         'hypothesis': 'Coverage gap or changed identifiers; cause is unconfirmed',
                         'next_step': 'Compare extraction filters, feed coverage, and key construction.'})
    fuzzy_info = {'enabled': False}
    detail = {'samples': samples, 'source_only_keys': heapq.nsmallest(50, (joined(k) for k in only_left)),
              'target_only_keys': heapq.nsmallest(50, (joined(k) for k in only_right))}
    if fuzzy is not None:
        threshold = Decimal(str(fuzzy))
        if not threshold.is_finite() or not 0 < threshold < 1:
            raise ValueError('Fuzzy threshold must be strictly between 0 and 1.')
        fuzzy_info = {'enabled': True, 'threshold': float(threshold), 'proposals': 0,
                      'note': 'Proposals are similarity leads to confirm manually; they are not matches and do not change counts.'}
        if only_left and only_right:
            if len(only_left) * len(only_right) > 4000000:
                fuzzy_info['skipped'] = 'Unmatched key volume is too large for fuzzy comparison; narrow the inputs or split the comparison.'
            else:
                proposals = []
                target_keys = [(b, joined(b), SequenceMatcher(None, '', joined(b)))
                                for b in sorted(only_right)]
                threshold_float = float(threshold)
                for a in sorted(only_left):
                    source_key = joined(a)
                    best, score = None, 0.0
                    for b, target_key, matcher in target_keys:
                        matcher.set_seq1(source_key)
                        if matcher.quick_ratio() < threshold_float:
                            continue
                        ratio = matcher.ratio()
                        if ratio >= threshold_float and ratio > score:
                            best, score = b, ratio
                    if best is not None:
                        proposals.append({'source_key': source_key, 'target_key': joined(best),
                                          'similarity': round(score, 4)})
                proposals.sort(key=lambda p: -p['similarity'])
                fuzzy_info['proposals'] = len(proposals)
                detail['fuzzy_proposals'] = proposals[:50]
                if proposals:
                    findings.append({'evidence': f'{len(proposals)} unmatched key pairs exceed the fuzzy threshold of {float(threshold)}',
                                     'hypothesis': 'Possible key typos or identifier variants; unconfirmed',
                                     'next_step': 'Confirm proposed pairs manually; never treat them as matched.'})
    schema = {'source_only': [f for f in lh if f not in rh], 'target_only': [f for f in rh if f not in lh]}
    total = shared_count + len(only_left) + len(only_right)
    rules = {'numeric_absolute_tolerance': str(tol), 'date_tolerance_days': str(days),
             'nulls': 'empty cells only; two empty cells match', 'strings': 'exact, case and whitespace sensitive'}
    if column_tolerances:
        rules['numeric_absolute_tolerance_per_column'] = {f: str(tols[f]) for f in column_tolerances}
    if column_days:
        rules['date_tolerance_days_per_column'] = {f: str(dayss[f]) for f in column_days}
    summary = {'version': 2, 'created_utc': datetime.now(timezone.utc).isoformat(),
               'key': keys[0] if len(keys) == 1 else keys,
               'source_rows': len(lr), 'target_rows': len(rr), 'shared_keys': shared_count,
               **dict(counts), 'source_only_rows': len(only_left), 'target_only_rows': len(only_right),
               'match_rate': (counts['exact_rows'] + counts['tolerance_rows']) / total if total else None,
               'schema_difference': schema, 'types': kinds,
               'rules': rules,
               'field_stats': {f: dict(c) for f, c in field_stats.items()},
               'cell_buckets': dict(buckets), 'findings': findings, 'fuzzy': fuzzy_info,
               'status': 'differences' if counts['mismatched_rows'] or only_left or only_right or any(schema.values()) else 'tied'}
    return summary, detail


def render(summary, detail):
    esc = lambda v: html.escape(str(v))
    cards = ''.join(f'<article><strong>{summary[k]}</strong><span>{label}</span></article>' for k, label in
                    [('exact_rows', 'Exact rows'), ('tolerance_rows', 'Within tolerance'),
                     ('mismatched_rows', 'Mismatched rows'), ('source_only_rows', 'Source only'), ('target_only_rows', 'Target only')])
    fields = ''.join(f'<tr><td>{esc(f)}</td><td>{esc(summary["types"][f])}</td><td>{esc(json.dumps(s))}</td></tr>' for f, s in summary['field_stats'].items())
    samples = ''.join('<tr>' + ''.join(f'<td>{esc(row[k])}</td>' for k in ('key', 'field', 'left', 'right', 'status', 'delta')) + '</tr>' for row in detail['samples'])
    findings = ''.join(f'<li><b>{esc(f["evidence"])}</b><p>{esc(f["hypothesis"])}</p>{esc(f["next_step"])}</li>' for f in summary['findings']) or '<li>No repeated offset detected. Inspect field counts and extraction rules for remaining differences.</li>'
    bars = ''.join(f'<div class="bar"><span>{esc(k)}: {v}</span><meter min="0" max="{max(summary["cell_buckets"].values(), default=1)}" value="{v}"></meter></div>' for k, v in summary['cell_buckets'].items())
    rate = 'N/A' if summary['match_rate'] is None else f'{summary["match_rate"]:.1%}'
    fuzzy = summary.get('fuzzy') or {}
    fuzzy_section = ''
    if fuzzy.get('enabled'):
        fuzzy_rows = ''.join('<tr>' + ''.join(f'<td>{esc(row[k])}</td>' for k in ('source_key', 'target_key', 'similarity')) + '</tr>'
                             for row in detail.get('fuzzy_proposals', []))
        fuzzy_body = (f'<div class="scroll"><table><tr><th>Source key</th><th>Target key</th><th>Similarity</th></tr>{fuzzy_rows}</table></div>'
                      if fuzzy_rows else '' if fuzzy.get('skipped') else '<p>No pairs above the threshold.</p>')
        skipped = f'<p>{esc(fuzzy["skipped"])}</p>' if fuzzy.get('skipped') else ''
        fuzzy_section = (f'<h2>Fuzzy key proposals</h2><p>Similarity leads for unmatched keys. '
                         f'Confirm manually; proposals are never automatic matches and do not change counts.</p>{skipped}{fuzzy_body}')
    return f'''<!doctype html><html lang="en"><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1">
<title>Tie-Out | Reconciliation report</title><style>
body{{font:16px system-ui,sans-serif;background:#f4f6f9;color:#18283b;margin:0}}main{{max-width:1150px;margin:auto;padding:40px 24px}}header{{border-top:5px solid #16817a;padding:24px 0}}h1{{font-size:46px;margin:8px 0}}h2{{margin-top:36px}}.eyebrow{{letter-spacing:3px;color:#16817a;font-size:12px;font-weight:700}}.cards{{display:flex;gap:12px;flex-wrap:wrap}}article{{background:white;border:1px solid #dce3eb;border-radius:12px;padding:20px;flex:1;min-width:120px}}strong{{display:block;font-size:32px}}article span{{color:#536379}}table{{border-collapse:collapse;width:100%;background:white}}td,th{{padding:12px;border-bottom:1px solid #dce3eb;text-align:left;overflow-wrap:anywhere}}.scroll{{overflow:auto}}pre{{white-space:pre-wrap;overflow-wrap:anywhere;background:white;padding:18px;border-radius:8px}}li{{margin:20px 0}}.bar{{display:flex;justify-content:space-between;gap:20px;margin:12px 0}}meter{{width:45%}}footer{{margin-top:40px;color:#536379;font-size:13px}}</style>
<main><header><div class="eyebrow">TIE-OUT / LOCAL DATA RECONCILIATION</div><h1>Do the numbers agree?</h1><p>Status: <b>{esc(summary['status'])}</b> · {rate} matching keys across the union · {summary['shared_keys']} shared keys</p></header>
<section class="cards">{cards}</section><h2>Discrepancy patterns</h2><p>Cell counts, not row counts. Tolerance matches are accepted differences, not errors.</p>{bars or '<p>No differing cells.</p>'}
<h2>Diagnosis & next steps</h2><p>Hypotheses describe possible causes, not proven root causes.</p><ul>{findings}</ul>
<h2>Field comparison</h2><div class="scroll"><table><tr><th>Field</th><th>Type</th><th>Cell counts</th></tr>{fields}</table></div>
<h2>Mismatch examples</h2><p>Up to 50 differing cells: largest absolute numeric and date differences first, remaining categories in key order. Magnitudes use different field units and do not establish business materiality; values are local report contents.</p><div class="scroll"><table><tr><th>Key</th><th>Field</th><th>Source</th><th>Target</th><th>Category</th><th>Delta</th></tr>{samples}</table></div>
<h2>Missing keys</h2><pre>{esc(json.dumps({k: v for k, v in detail.items() if k not in ('samples', 'fuzzy_proposals')}, indent=2))}</pre>
{fuzzy_section}<h2>Comparison contract</h2><pre>{esc(json.dumps({'key':summary['key'], **summary['rules'], 'schema_difference':summary['schema_difference']}, indent=2))}</pre>
<footer>Generated {esc(summary['created_utc'])}. Match rate = (exact + within-tolerance rows) / union of keys. Empty inputs have no match rate. Schema differences prevent a tied status. Fuzzy proposals never affect counts. No data was uploaded by this script.</footer></main></html>'''


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('mode', choices=['profile', 'compare'])
    parser.add_argument('source', nargs='?', help='Input file path, or a placeholder such as - when --source-query is used')
    parser.add_argument('target', nargs='?', help='Input file path, or a placeholder such as - when --target-query is used')
    parser.add_argument('--key', help='Join key column, or comma-separated columns for a composite key')
    parser.add_argument('--source-sheet')
    parser.add_argument('--target-sheet')
    parser.add_argument('--source-query', help='Read the source side from SQL Server with this read-only SELECT/WITH query')
    parser.add_argument('--target-query', help='Read the target side from SQL Server with this read-only SELECT/WITH query')
    parser.add_argument('--mssql-dsn', help='SQL Server ODBC connection string; defaults to the TIEOUT_MSSQL_DSN environment variable. Never logged.')
    parser.add_argument('--max-rows', type=int, default=1000000, help='Reject SQL results above this row count instead of silently truncating')
    parser.add_argument('--timeout', type=int, default=60, help='SQL Server login and query timeout in seconds')
    parser.add_argument('--numeric-tolerance', default='0.01')
    parser.add_argument('--date-days', default='0')
    parser.add_argument('--tolerance', action='append', default=[], metavar='COLUMN=VALUE', help='Per-column numeric absolute tolerance; overrides --numeric-tolerance for that column')
    parser.add_argument('--date-tolerance', action='append', default=[], metavar='COLUMN=DAYS', help='Per-column date tolerance in days; overrides --date-days for that column')
    parser.add_argument('--type', action='append', default=[], metavar='COLUMN=TYPE')
    parser.add_argument('--fuzzy', help='Report similarity proposals for unmatched keys above this threshold, strictly between 0 and 1; proposals are diagnostic and never count as matches')
    parser.add_argument('--out', default='reports/tie-out')
    parser.add_argument('--fail-on-difference', action='store_true')
    args = parser.parse_args(argv)
    try:
        dsn = args.mssql_dsn or os.environ.get('TIEOUT_MSSQL_DSN')

        def side(path, query, sheet):
            if query is not None:
                if sheet is not None:
                    raise ValueError('Sheet selection applies to XLSX files, not queries.')
                return load_sql(query, dsn, args.max_rows, args.timeout)
            if path is None:
                raise ValueError('Each side needs an input file or a --source-query / --target-query.')
            return load(path, sheet)

        left, right = side(args.source, args.source_query, args.source_sheet), side(args.target, args.target_query, args.target_sheet)
        if args.mode == 'profile':
            lp, rp = profile(*left), profile(*right)
            candidates = [f for f in left[0] if f in right[0] and lp['columns'][f]['candidate_key'] and rp['columns'][f]['candidate_key']]
            print(json.dumps({'source': lp, 'target': rp, 'candidate_keys': candidates}, indent=2))
            return 0
        if not args.key:
            raise ValueError('Compare requires --key; run profile to inspect candidate keys.')
        overrides = dict(item.split('=', 1) for item in args.type)
        column_tolerances = dict(item.split('=', 1) for item in args.tolerance)
        column_days = dict(item.split('=', 1) for item in args.date_tolerance)
        summary, detail = reconcile(left, right, args.key, args.numeric_tolerance, args.date_days, overrides,
                                    column_tolerances, column_days, args.fuzzy)
        out = Path(args.out)
        targets = [out / 'summary.json', out / 'report.html']
        input_paths = [Path(p).resolve() for p in (args.source, args.target) if p]
        if any(p.resolve() in input_paths for p in targets):
            raise ValueError('Output must not overwrite an input.')
        if any(p.exists() for p in targets):
            raise ValueError('Output already exists; choose a new --out directory.')
        out.mkdir(parents=True, exist_ok=True)
        targets[0].write_text(json.dumps(summary, indent=2), encoding='utf-8')
        targets[1].write_text(render(summary, detail), encoding='utf-8')
        print(json.dumps({'summary': summary, 'report': str(targets[1].resolve())}, indent=2))
        return 1 if args.fail_on_difference and summary['status'] == 'differences' else 0
    except (ValueError, OSError, csv.Error, ImportError, InvalidOperation, BadZipFile) as error:
        print(f'tie-out: {error}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    sys.exit(main())
